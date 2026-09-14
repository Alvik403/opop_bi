from __future__ import annotations

import pytest
from openpyxl import load_workbook
from zipfile import ZIP_DEFLATED, ZipFile

from file_validation import ValidationLimits, validate_excel_file


def test_validate_excel_file_accepts_expected_structure(sample_excel_path):
    result = validate_excel_file(sample_excel_path)

    assert result.service_count == 1
    assert result.class_count == 1


def test_validate_excel_file_rejects_non_xlsx(tmp_path):
    path = tmp_path / "data.txt"
    path.write_text("not excel", encoding="utf-8")

    with pytest.raises(ValueError, match="только файлы .xlsx"):
        validate_excel_file(path)


def test_validate_excel_file_rejects_wrong_structure(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an excel file")

    with pytest.raises(ValueError, match="ZIP-контейнер"):
        validate_excel_file(path)


def test_validate_excel_file_rejects_unsafe_compression_ratio(tmp_path):
    path = tmp_path / "bomb.xlsx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"0" * 100_000)
        archive.writestr("xl/workbook.xml", b"<workbook />")

    with pytest.raises(ValueError, match="коэффициент сжатия"):
        validate_excel_file(path, ValidationLimits(max_compression_ratio=2))


def test_validate_excel_file_rejects_too_many_sheets(sample_excel_path):
    workbook = load_workbook(sample_excel_path)
    workbook.create_sheet("Лишний лист")
    workbook.save(sample_excel_path)
    workbook.close()

    with pytest.raises(ValueError, match="количества листов"):
        validate_excel_file(sample_excel_path, ValidationLimits(max_sheets=1))


def test_validate_excel_file_rejects_too_many_rows(sample_excel_path):
    workbook = load_workbook(sample_excel_path)
    workbook["Калькуляция"].cell(row=101, column=1, value="overflow")
    workbook.save(sample_excel_path)
    workbook.close()

    with pytest.raises(ValueError, match="лимит строк"):
        validate_excel_file(sample_excel_path, ValidationLimits(max_rows_per_sheet=100))
