from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from starlette.testclient import TestClient

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

APP_MODULES = [
    "app",
    "settings",
    "file_registry",
    "file_security",
    "dashboard_cache",
    "logging_config",
    "security",
]


def create_sample_workbook(
    path: Path,
    *,
    class_name: str = "экспорт/сибур",
    service_name: str = "Комплексная/услуга 2",
    sheet_title: str = "Калькуляция",
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title

    top_headers = [
        "№",
        "Класс",
        "Услуга",
        "Комментарий",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Прямые расходы",
        "Косвенные расходы",
        "Косвенные расходы",
        "Косвенные расходы",
        "Косвенные расходы",
        "Косвенные расходы",
        "Косвенные расходы",
        "Неэффективность",
        "Неэффективность",
        "Неэффективность",
    ]
    sub_headers = [
        "",
        "",
        "",
        "",
        "ИТОГО",
        "ФОТ",
        "Материалы",
        "Процент %",
        "Прочее",
        "Деталь 5",
        "Деталь 6",
        "Деталь 7",
        "Деталь 8",
        "Деталь 9",
        "Деталь 10",
        "Деталь 11",
        "ИТОГО",
        "Аренда",
        "Связь",
        "Охрана",
        "Прочее",
        "Деталь косвенная",
        "Аренда",
        "РЖД",
        "ПРТ",
    ]
    values = [
        1,
        class_name,
        service_name,
        "",
        1000,
        600,
        400,
        10,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        0,
        500,
        200,
        100,
        50,
        50,
        100,
        100,
        50,
        25,
    ]

    for col_idx, value in enumerate(top_headers, start=1):
        sheet.cell(row=3, column=col_idx, value=value)
    for col_idx, value in enumerate(sub_headers, start=1):
        sheet.cell(row=4, column=col_idx, value=value)
    for col_idx, value in enumerate(values, start=1):
        sheet.cell(row=5, column=col_idx, value=value)

    workbook.save(path)


def load_app(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, **env: str) -> TestClient:
    data_dir = tmp_path / "data"
    if not (data_dir / "data.xlsx").exists():
        create_sample_workbook(data_dir / "data.xlsx")

    defaults = {
        "DEBUG": "false",
        "SESSION_SECRET": "x" * 32,
        "AUTH_USERNAME": "admin",
        "AUTH_PASSWORD": "secret-password-123",
        "DATA_DIR": str(data_dir),
        "UPLOADS_DIR": str(tmp_path / "uploads"),
        "DATABASE_PATH": str(tmp_path / "runtime" / "test.sqlite3"),
        "LOGS_DIR": str(tmp_path / "logs"),
        "MAX_UPLOAD_BYTES": "1024",
        "UPLOAD_RATE_LIMIT_PER_MINUTE": "2",
        "ALLOWED_HOSTS": "testserver",
        "AUTH_LOCKOUT_ATTEMPTS": "5",
        "AUTH_LOCKOUT_WINDOW_SECONDS": "300",
    }
    defaults.update(env)
    for key, value in defaults.items():
        if value is None:
            monkeypatch.delenv(key, raising=False)
        else:
            monkeypatch.setenv(key, value)

    for module_name in APP_MODULES:
        sys.modules.pop(module_name, None)

    spec = importlib.util.spec_from_file_location("app", PROJECT_ROOT / "app.py")
    assert spec and spec.loader
    app_module = importlib.util.module_from_spec(spec)
    sys.modules["app"] = app_module
    spec.loader.exec_module(app_module)
    base_url = "http://testserver" if defaults.get("DEBUG") == "true" else "https://testserver"
    return TestClient(app_module.app, base_url=base_url)


def csrf_headers(client: TestClient, extra: dict[str, str] | None = None) -> dict[str, str]:
    token = client.get("/api/csrf-token", headers=extra or {}).json()["csrf_token"]
    origin = str(client.base_url).rstrip("/")
    headers = {"X-CSRF-Token": token, "Origin": origin}
    if extra:
        headers.update(extra)
    return headers


@pytest.fixture()
def sample_excel_path(tmp_path: Path) -> Path:
    path = tmp_path / "data.xlsx"
    create_sample_workbook(path)
    return path


@pytest.fixture()
def app_client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    with load_app(
        monkeypatch,
        tmp_path,
        DEBUG="true",
        ALLOW_INSECURE_DEBUG="true",
        AUTH_USERNAME="",
        AUTH_PASSWORD="",
        SESSION_SECRET="test-secret",
        MAX_UPLOAD_BYTES="52428800",
        UPLOAD_RATE_LIMIT_PER_MINUTE="10",
    ) as client:
        yield client
