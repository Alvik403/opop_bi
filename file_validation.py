from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from data_loader import load_calculation_services_dataset


@dataclass(frozen=True)
class ValidationLimits:
    max_archive_entries: int = 20_000
    max_uncompressed_bytes: int = 262_144_000
    max_compression_ratio: int = 100
    max_sheets: int = 50
    max_rows_per_sheet: int = 100_000


@dataclass(frozen=True)
class ValidationResult:
    service_count: int
    class_count: int

    def to_dict(self) -> dict:
        return {
            "service_count": self.service_count,
            "class_count": self.class_count,
        }


def _validate_xlsx_container(path: Path, limits: ValidationLimits) -> None:
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_archive_entries:
                raise ValueError("В XLSX слишком много элементов архива")

            names = {entry.filename for entry in entries}
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names):
                raise ValueError("Файл не является корректной книгой XLSX")

            total_uncompressed = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise ValueError("Зашифрованные XLSX не поддерживаются")
                total_uncompressed += entry.file_size
                if total_uncompressed > limits.max_uncompressed_bytes:
                    raise ValueError("Распакованный XLSX превышает допустимый объём")
                if entry.file_size and (
                    entry.compress_size == 0
                    or entry.file_size / entry.compress_size > limits.max_compression_ratio
                ):
                    raise ValueError("XLSX имеет небезопасный коэффициент сжатия")
    except BadZipFile as exc:
        raise ValueError("Файл не является корректным ZIP-контейнером XLSX") from exc


def _validate_workbook_dimensions(path: Path, limits: ValidationLimits) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise ValueError("В XLSX превышен лимит количества листов")
        for sheet in workbook.worksheets:
            if sheet.max_row > limits.max_rows_per_sheet:
                raise ValueError(
                    f"На листе {sheet.title!r} превышен лимит строк"
                )
    finally:
        workbook.close()


def validate_excel_file(
    path: Path,
    limits: ValidationLimits | None = None,
) -> ValidationResult:
    limits = limits or ValidationLimits()
    if path.suffix.casefold() != ".xlsx":
        raise ValueError("Поддерживаются только файлы .xlsx")
    if not path.is_file():
        raise ValueError("Файл XLSX не найден")

    try:
        _validate_xlsx_container(path, limits)
        _validate_workbook_dimensions(path, limits)
        dataset = load_calculation_services_dataset(path)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Файл не прошёл проверку структуры") from exc

    services = dataset.get("services", [])
    classes = {service.get("class_name") for service in services if service.get("class_name")}
    if not services:
        raise ValueError("В файле не найдены услуги на листе 'Калькуляция'")
    if not classes:
        raise ValueError("В файле не найдены классы услуг")

    required_detail_groups = (
        "direct_detail_labels",
        "indirect_detail_labels",
        "inefficiency_detail_labels",
    )
    missing_groups = [name for name in required_detail_groups if not dataset.get(name)]
    if missing_groups:
        raise ValueError(f"Не найдены ожидаемые группы стоимостных колонок: {', '.join(missing_groups)}")

    return ValidationResult(service_count=len(services), class_count=len(classes))
